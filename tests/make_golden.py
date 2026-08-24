"""Trích fixture golden từ file Excel gốc. Chạy một lần, kết quả commit vào repo.

    python tests/make_golden.py "D:/BP/Bros Partners/Tickers/VHC/Important file VHC/giá cá nhập khẩu US - (final).xlsx"

Input golden lấy từ Sheet1 (dữ liệu thô), expected lấy từ Sheet2 (giá trị đã
tính sẵn trong Excel). build.py phải tự đi từ cái thứ nhất tới cái thứ hai.
"""

import csv
import json
import os
import sys
from collections import defaultdict

import openpyxl

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from scripts.console import fix_stdio_encoding

fix_stdio_encoding()

MONTHS = ["January", "February", "March", "April", "May", "June", "July",
          "August", "September", "October", "November", "December"]
MONTH_NUM = {name: f"{i + 1:02d}" for i, name in enumerate(MONTHS)}

FIXTURES = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fixtures")

# Bố cục Sheet2, bóc từ công thức trong file gốc.
LAYOUT = [
    {"key": "tilapia", "name_cell": 6, "volume": 7, "value": 8, "asp": 30,
     "breakdown_volume": {"CHINA": 13, "VIETNAM": 14, "TAIWAN": 15,
                          "INDONESIA": 16},
     "breakdown_value": {"CHINA": 19, "VIETNAM": 20, "TAIWAN": 21,
                         "INDONESIA": 22},
     "breakdown_asp": {"CHINA": 31, "TAIWAN": 33, "INDONESIA": 34}},
    {"key": "pangasius", "name_cell": 49, "volume": 50, "value": 51, "asp": 59,
     "breakdown_volume": {"VIETNAM": 56, "TAIWAN": 57},
     "breakdown_value": {}, "breakdown_asp": {}},
    {"key": "haddock", "name_cell": 63, "volume": 64, "value": 65, "asp": 78,
     "breakdown_volume": {"CHINA": 71, "INDONESIA": 72, "NORWAY": 73,
                          "CANADA": 74, "ICELAND": 75},
     "breakdown_value": {}, "breakdown_asp": {}},
    {"key": "salmon", "name_cell": 82, "volume": 83, "value": 84, "asp": 87,
     "breakdown_volume": {}, "breakdown_value": {}, "breakdown_asp": {}},
    {"key": "cod", "name_cell": 91, "volume": 92, "value": 93, "asp": 107,
     "breakdown_volume": {"CHINA": 98, "INDONESIA": 99, "NORWAY": 100,
                          "CANADA": 101, "ICELAND": 102, "ECUADOR": 103,
                          "GREENLAND": 104, "VIETNAM": 105},
     "breakdown_value": {}, "breakdown_asp": {}},
    {"key": "pollock", "name_cell": 110, "volume": 111, "value": 112,
     "asp": 126,
     "breakdown_volume": {"CHINA": 117, "INDONESIA": 118,
                          "UNITED KINGDOM": 119, "CANADA": 120,
                          "ICELAND": 121, "ECUADOR": 122, "GREENLAND": 123,
                          "VIETNAM": 124},
     "breakdown_value": {}, "breakdown_asp": {}},
]


def period_columns(sheet2):
    """Đọc hàng 3 (năm) và hàng 5 (tên tháng) -> {"2023-01": col_index}."""
    periods = {}
    for col in range(4, 60):
        year = sheet2.cell(3, col).value
        month = sheet2.cell(5, col).value
        if year is None or month is None:
            continue
        periods[f"{int(year)}-{MONTH_NUM[month]}"] = col
    return periods


def extract_sheet1(sheet1, products):
    """Gộp Sheet1 về grain (year, month, product, country) cho các SP quan tâm."""
    agg = defaultdict(lambda: [0, 0])
    for row in sheet1.iter_rows(min_row=3, values_only=True):
        if not row[0] or row[3] not in products:
            continue
        key = (str(row[0]), MONTH_NUM[row[1]], row[3], row[4])
        agg[key][0] += row[5] or 0
        agg[key][1] += row[6] or 0
    return agg


def cell(sheet2, row, col):
    value = sheet2.cell(row, col).value
    if isinstance(value, str):  # ô lỗi kiểu "#DIV/0!"
        return None
    return value


def main():
    path = sys.argv[1]
    book = openpyxl.load_workbook(path, data_only=True)
    sheet1, sheet2 = book["Sheet1"], book["Sheet2"]

    for spec in LAYOUT:
        spec["product"] = sheet2.cell(spec["name_cell"], 3).value

    products = {spec["product"] for spec in LAYOUT}
    periods = period_columns(sheet2)
    ordered = sorted(periods)

    os.makedirs(FIXTURES, exist_ok=True)

    agg = extract_sheet1(sheet1, products)
    csv_path = os.path.join(FIXTURES, "sheet1_rows.csv")
    with open(csv_path, "w", newline="", encoding="utf8") as fh:
        writer = csv.writer(fh, lineterminator="\n")
        writer.writerow(["year", "month", "product", "country",
                         "volume_kg", "value_usd"])
        for (year, month, product, country), (volume, value) in sorted(agg.items()):
            if f"{year}-{month}" in periods:
                writer.writerow([year, month, product, country, volume, value])

    expected = {"months": ordered, "groups": {}}
    for spec in LAYOUT:
        entry = {
            "product": spec["product"],
            "volume": [cell(sheet2, spec["volume"], periods[p]) for p in ordered],
            "value": [cell(sheet2, spec["value"], periods[p]) for p in ordered],
            "asp": [cell(sheet2, spec["asp"], periods[p]) for p in ordered],
            "country_volume": {}, "country_value": {}, "country_asp": {},
        }
        for field, name in (("breakdown_volume", "country_volume"),
                            ("breakdown_value", "country_value"),
                            ("breakdown_asp", "country_asp")):
            for country, row in spec[field].items():
                entry[name][country] = [cell(sheet2, row, periods[p])
                                        for p in ordered]
        expected["groups"][spec["key"]] = entry

    json_path = os.path.join(FIXTURES, "sheet2_expected.json")
    with open(json_path, "w", encoding="utf8") as fh:
        json.dump(expected, fh, ensure_ascii=False, indent=1)

    print(f"{csv_path}: {sum(1 for _ in open(csv_path, encoding='utf8')) - 1} dòng")
    print(f"{json_path}: {len(ordered)} tháng, {len(expected['groups'])} nhóm")
    print(f"Khoảng thời gian: {ordered[0]} -> {ordered[-1]}")


if __name__ == "__main__":
    main()
